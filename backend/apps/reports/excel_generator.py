from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
from datetime import datetime


def generate_report_excel(title, headers, rows, summary=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal='center')

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(row=2, column=1, value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    # Headers
    header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font

    # Rows
    for r, row in enumerate(rows, 5):
        for c, val in enumerate(row, 1):
            ws.cell(row=r, column=c, value=val)

    # Column widths
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 18

    # Summary sheet
    if summary:
        ws2 = wb.create_sheet("Summary")
        ws2.cell(row=1, column=1, value="Metric").font = Font(bold=True)
        ws2.cell(row=1, column=2, value="Value").font = Font(bold=True)
        for i, (k, v) in enumerate(summary.items(), 2):
            ws2.cell(row=i, column=1, value=k)
            ws2.cell(row=i, column=2, value=str(v))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer